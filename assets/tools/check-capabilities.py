import tempfile
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from reportlab.pdfgen.canvas import Canvas
from pypdf import PdfReader
from playwright.sync_api import sync_playwright


with tempfile.TemporaryDirectory() as directory:
    root = Path(directory)
    document = Document()
    document.add_paragraph('ok')
    document.save(root / 'check.docx')
    workbook = Workbook()
    workbook.active['A1'] = 'ok'
    workbook.save(root / 'check.xlsx')
    Presentation().save(root / 'check.pptx')
    canvas = Canvas(str(root / 'check.pdf'))
    canvas.drawString(72, 720, 'ok')
    canvas.save()
    assert len(PdfReader(root / 'check.pdf').pages) == 1
    assert all((root / name).stat().st_size > 0 for name in ('check.docx', 'check.xlsx', 'check.pptx', 'check.pdf'))

with sync_playwright() as playwright:
    browser = playwright.chromium.launch(channel='msedge', headless=True)
    page = browser.new_page()
    page.goto('about:blank')
    assert page.url == 'about:blank'
    browser.close()

print('documents-and-browser-ok')
